"""Compare two MC12 backtest reports before you compare their numbers.

WHY THIS EXISTS
---------------
2026-08-23. Two runs of the same L1 v32 file, same 42 settings fields, produced
500 trades / 5,483,600 and 501 trades / 5,388,800. It looked like MC12 was
non-deterministic, which would have invalidated every optimisation decision
ever taken on this platform.

It was not. The first 500 trades were identical field for field. The second run
simply had 9 hours 15 minutes more data and completed one extra trade worth
exactly the -94,800 difference.

The trap: BOTH reports state an end date of 2026-08-22 05:00 in the settings
sheet. That field is the REQUESTED range, not the last bar the run actually
had. The only place the real data end shows up is the exit timestamp of the
last closed trade.

So: never compare two MC12 reports without first checking they saw the same
data. This script does that check, then localises any divergence.

USAGE
-----
  python scripts/compare_mc12_reports.py A.xlsx B.xlsx

EXIT CODES
----------
  0  comparable and identical
  1  comparable but the trades diverge  -> a real behaviour difference
  2  NOT comparable (different data windows or different settings)
     -> fix that before reading a single performance number
"""
import sys

try:
    import openpyxl
except ImportError:
    sys.exit('openpyxl required: pip install openpyxl')

import datetime

TRADES = '交易明細'   # trade detail
SETUP = '設定'                # settings
ENTER = '進入'                # "enter"
LEAVE = '離開'                # "leave"

# Settings rows that are metadata about WHEN the run happened rather than WHAT
# it ran, so a difference in them is not by itself an incomparability.
IGNORE = set()


def read_settings(path):
    """Return (strategy_name, {field: value}).

    The settings sheet opens with a sheet-title row and then a row carrying
    the strategy name with no value. Those two are identity, not settings.
    Leaving them in the dict made every cross-version compare report a
    settings mismatch on the strategy name alone.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    if SETUP not in wb.sheetnames:
        sys.exit('%s has no settings sheet' % path)
    rows = [r for r in wb[SETUP].iter_rows(values_only=True)
            if r and r[0] is not None]
    name = ''
    out = {}
    for i, r in enumerate(rows):
        k = str(r[0]).strip()
        v = '' if len(r) < 2 or r[1] is None else str(r[1]).strip()
        if i < 2 and v == '':
            if k != SETUP:
                name = k
            continue
        out[k] = v
    return name, out


def read_trades(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    if TRADES not in wb.sheetnames:
        sys.exit('%s has no trade detail sheet' % path)
    ws = wb[TRADES]
    rows = []
    for r in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        if any(v is not None for v in r[:4]):
            rows.append(tuple('' if v is None else str(v).strip() for v in r))

    def dt(d, t):
        return datetime.datetime.strptime(d[:10] + ' ' + t, '%Y-%m-%d %H:%M:%S')

    out = []
    i = 0
    while i < len(rows) - 1:
        e, x = rows[i], rows[i + 1]
        if e[2].startswith(ENTER) and x[2].startswith(LEAVE):
            out.append({
                'entry_sig': e[3], 'exit_sig': x[3],
                't_in': dt(e[4], e[5]), 't_out': dt(x[4], x[5]),
                'p_in': float(e[6]), 'p_out': float(x[6]),
                'pnl': float(e[8] or 0),
                'mfe': float(e[12] or 0), 'mae': float(e[14] or 0),
            })
            i += 2
        else:
            i += 1
    if not out:
        sys.exit('%s: parsed zero trades' % path)
    return out


def key(t):
    return (t['entry_sig'], t['exit_sig'], t['t_in'], t['t_out'],
            round(t['p_in'], 4), round(t['p_out'], 4), round(t['pnl'], 4))


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    pa, pb = sys.argv[1], sys.argv[2]
    na, sa = read_settings(pa)
    nb, sb = read_settings(pb)
    ta, tb = read_trades(pa), read_trades(pb)

    print('=' * 74)
    print('A: %s' % pa)
    print('B: %s' % pb)
    print('=' * 74)

    verdict = 0

    # ---- gate 1: settings ----
    # A field present in only ONE file is a VERSION DELTA, not an
    # incomparability. v17 adding four form switches does not by itself mean
    # the two runs differ -- whether the new fields changed anything is
    # exactly what gate 3 measures, and gate 3 is the authority.
    # Only fields present in BOTH files with different values are a genuine
    # mismatch. (2026-08-24: the first version of this failed the L4 v17
    # anchor with EXIT=2 while all 79 trades were identical.)
    shared = sorted(set(sa) & set(sb))
    only_a = sorted(set(sa) - set(sb))
    only_b = sorted(set(sb) - set(sa))
    sdiff = [(k, sa[k], sb[k]) for k in shared
             if k not in IGNORE and sa[k] != sb[k]]
    print('\n[1] SETTINGS   A=%s   B=%s' % (na or '?', nb or '?'))
    print('      %d shared fields, %d differ' % (len(shared), len(sdiff)))
    for k, x, y in sdiff:
        print('      ** %-26s A=%-24s B=%s' % (k, x, y))
    if only_a or only_b:
        print('      version delta -- informational, gate 3 decides:')
        for k in only_a:
            print('         only in A   %-24s = %s' % (k, sa[k]))
        for k in only_b:
            print('         only in B   %-24s = %s' % (k, sb[k]))
    if sdiff:
        verdict = 2

    # ---- gate 2: the data window each run ACTUALLY saw ----
    # The settings end date is the requested range. The real end is the last
    # closed-trade exit. See module docstring.
    a_end, b_end = ta[-1]['t_out'], tb[-1]['t_out']
    a_beg, b_beg = ta[0]['t_in'], tb[0]['t_in']
    print('\n[2] DATA WINDOW ACTUALLY USED  (last closed exit, not the settings field)')
    print('      A  %s  ..  %s' % (a_beg, a_end))
    print('      B  %s  ..  %s' % (b_beg, b_end))
    if a_beg != b_beg or a_end != b_end:
        gap = abs((b_end - a_end).total_seconds()) / 3600.0
        print('      ** DIFFERENT. tail gap %.2f hours **' % gap)
        print('      ** These two runs did NOT see the same data. **')
        print('      ** Any performance comparison between them is invalid. **')
        verdict = 2
    else:
        print('      identical')

    # ---- gate 3: trade-by-trade ----
    n = min(len(ta), len(tb))
    first = None
    for i in range(n):
        if key(ta[i]) != key(tb[i]):
            first = i
            break
    print('\n[3] TRADES  A=%d  B=%d  common prefix compared=%d' % (len(ta), len(tb), n))
    if first is None:
        print('      first %d trades identical field for field' % n)
        if len(ta) != len(tb):
            longer, extra = ('B', tb[n:]) if len(tb) > len(ta) else ('A', ta[n:])
            tot = sum(t['pnl'] for t in extra)
            print('      %s has %d extra trade(s) at the tail, worth %s'
                  % (longer, len(extra), format(int(tot), ',')))
            for t in extra:
                print('        %s @%.0f -> %s @%.0f  %-14s %s'
                      % (t['t_in'], t['p_in'], t['t_out'], t['p_out'],
                         t['exit_sig'], format(int(t['pnl']), ',')))
            na, nb = sum(t['pnl'] for t in ta), sum(t['pnl'] for t in tb)
            print('      net A %s  net B %s  diff %s'
                  % (format(int(na), ','), format(int(nb), ','),
                     format(int(nb - na), ',')))
            if abs(abs(nb - na) - abs(tot)) < 0.01:
                print('      ** the whole difference IS the tail trade(s) --'
                      ' behaviour is identical, only the data length differs **')
            if verdict == 0:
                verdict = 2
    else:
        print('      ** FIRST DIVERGENCE at trade #%d **' % (first + 1))
        for lbl, S in (('A', ta), ('B', tb)):
            t = S[first]
            print('        %s  %s @%.0f -> %s @%.0f  %-14s %s'
                  % (lbl, t['t_in'], t['p_in'], t['t_out'], t['p_out'],
                     t['exit_sig'], format(int(t['pnl']), ',')))
        if verdict == 0:
            verdict = 1

    print('\n' + '=' * 74)
    print({0: 'VERDICT: comparable and identical',
           1: 'VERDICT: comparable, but the trades diverge -- real behaviour difference',
           2: 'VERDICT: NOT COMPARABLE -- fix this before reading any number'}[verdict])
    if verdict == 0 and (only_a or only_b):
        print('DEGENERACY ANCHOR HELD -- the new fields changed nothing.')
    print('=' * 74)
    return verdict


if __name__ == '__main__':
    sys.exit(main())
