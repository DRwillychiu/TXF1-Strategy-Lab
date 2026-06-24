"""S3_S A/B Cooldown 0 vs 1 detailed comparison."""
import sys, io
from collections import defaultdict
from openpyxl import load_workbook
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

PATH_A = r'C:\Users\User\Downloads\TXF1  VolSqueezeShort 策略回測績效報告.xlsx'  # Cool=0
PATH_B = r'C:\Users\User\Downloads\TXF1  VolSqueezeShort 策略回測績效報告_COOL1.xlsx'  # Cool=1

def extract(path):
    wb = load_workbook(path, data_only=True)
    out = {'metrics': {}, 'trades': [], 'inputs': {}}
    # Performance
    ws = wb['策略分析']
    for row in ws.iter_rows(min_row=1, max_row=100, values_only=True):
        if not row[0]: continue
        label = str(row[0]).strip()
        if label in ['淨利', '毛利', '毛損', '獲利因子', '調整獲利因子', '滑價支付',
                     '交易總次數', '%勝率', '年度夏普比率', '索丁諾比率',
                     '年報酬率', '最大策略虧損', '最大策略虧損 (%)',
                     '最大連續虧損', '平均交易獲利', '最大單筆獲利', '最大單筆虧損']:
            out['metrics'][label] = row[1]
    # Inputs
    ws = wb['設定']
    for row in ws.iter_rows(values_only=True):
        if row[0] and row[1] is not None:
            out['inputs'][str(row[0])] = row[1]
    # Trades
    ws = wb['交易明細']
    pending = None
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or not row[3]: continue
        sig = str(row[3]) if row[3] else None
        dt = row[4]
        pnl = row[8] if len(row) > 8 else None
        if sig and 'SE_' in sig:
            entry_d = dt.date() if hasattr(dt, 'date') else dt
            pnl_val = pnl if isinstance(pnl, (int, float)) else 0
            pending = {'entry_d': entry_d, 'pnl': pnl_val, 'entry_sig': sig}
        elif sig and 'SX_' in sig and pending is not None:
            exit_d = dt.date() if hasattr(dt, 'date') else dt
            out['trades'].append({**pending, 'exit_d': exit_d, 'exit_sig': sig})
            pending = None
    return out

A = extract(PATH_A)  # Cool=0
B = extract(PATH_B)  # Cool=1

# Confirm inputs differ only on Cooldown_Days
print('=' * 80)
print('INPUT VERIFICATION (should differ only on Cooldown_Days)')
print('=' * 80)
key_inputs = ['BBLen','BWPctile','StopATRMult','TargetATRMult','MaxBars',
              'MidExit_MinBars','SP_Trigger_ATRMult','SP_Retain_Pct','Cooldown_Days']
print(f'{"Input":<22} {"Cool=0":>12} {"Cool=1":>12}')
for k in key_inputs:
    a = A['inputs'].get(k, '?')
    b = B['inputs'].get(k, '?')
    marker = '  *** DIFFER ***' if a != b else ''
    print(f'  {k:<20} {str(a):>12} {str(b):>12}{marker}')

# Performance side-by-side
print()
print('=' * 80)
print('PERFORMANCE COMPARISON')
print('=' * 80)
print(f'{"Metric":<22} {"Cool=0":>14} {"Cool=1":>14} {"Diff":>14} {"%":>8}')
print('-' * 80)
for key in ['淨利', '毛利', '毛損', '獲利因子', '調整獲利因子', '滑價支付',
            '交易總次數', '%勝率', '年度夏普比率', '索丁諾比率',
            '年報酬率', '最大策略虧損', '最大策略虧損 (%)',
            '平均交易獲利', '最大單筆獲利', '最大單筆虧損']:
    a = A['metrics'].get(key, 0) or 0
    b = B['metrics'].get(key, 0) or 0
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        diff = b - a
        pct = (diff / abs(a) * 100) if a != 0 else 0
        better = ''
        # heuristic better-direction
        if key in ['淨利','毛利','獲利因子','調整獲利因子','%勝率',
                   '年度夏普比率','索丁諾比率','年報酬率',
                   '平均交易獲利','最大單筆獲利']:
            better = ' ✓' if diff > 0 else (' ✗' if diff < 0 else '')
        elif key in ['毛損','滑價支付','最大策略虧損','最大策略虧損 (%)',
                     '最大單筆虧損']:
            # less negative / smaller absolute is better (these are negative or cost)
            better = ' ✓' if abs(b) < abs(a) else (' ✗' if abs(b) > abs(a) else '')
        if isinstance(a, float) and abs(a) < 100:
            print(f'  {key:<20} {a:>14.3f} {b:>14.3f} {diff:>+14.3f} {pct:>+7.1f}%{better}')
        else:
            print(f'  {key:<20} {a:>14,.0f} {b:>14,.0f} {diff:>+14,.0f} {pct:>+7.1f}%{better}')

# Year distribution
print()
print('=' * 80)
print('YEAR-BY-YEAR COMPARISON')
print('=' * 80)
def year_stats(trades):
    s = defaultdict(lambda: {'n':0, 'pnl':0, 'wins':0})
    for t in trades:
        y = t['entry_d'].year
        s[y]['n'] += 1
        s[y]['pnl'] += t['pnl']
        if t['pnl'] > 0: s[y]['wins'] += 1
    return s
ya, yb = year_stats(A['trades']), year_stats(B['trades'])
all_years = sorted(set(list(ya.keys()) + list(yb.keys())))
print(f'{"Year":<6} {"Cool=0 N/PnL":>22} {"Cool=1 N/PnL":>22} {"N diff":>8} {"PnL diff":>14}')
print('-' * 80)
for y in all_years:
    a = ya[y]; b = yb[y]
    print(f'{y:<6} {a["n"]:>6} / {a["pnl"]:>+12,.0f} {b["n"]:>6} / {b["pnl"]:>+12,.0f} '
          f'{b["n"]-a["n"]:>+8} {b["pnl"]-a["pnl"]:>+14,.0f}')

# Exit distribution
print()
print('=' * 80)
print('EXIT SIGNAL DISTRIBUTION COMPARISON')
print('=' * 80)
def exit_stats(trades):
    s = defaultdict(lambda: {'n':0, 'pnl':0, 'wins':0})
    for t in trades:
        s[t['exit_sig']]['n'] += 1
        s[t['exit_sig']]['pnl'] += t['pnl']
        if t['pnl'] > 0: s[t['exit_sig']]['wins'] += 1
    return s
ea, eb = exit_stats(A['trades']), exit_stats(B['trades'])
all_sigs = sorted(set(list(ea.keys()) + list(eb.keys())))
print(f'{"ExitSig":<22} {"Cool=0 N/WR/PnL":>26} {"Cool=1 N/WR/PnL":>26}')
print('-' * 80)
for sig in all_sigs:
    a = ea[sig]; b = eb[sig]
    wra = a['wins']/a['n']*100 if a['n'] else 0
    wrb = b['wins']/b['n']*100 if b['n'] else 0
    print(f'  {sig:<20} {a["n"]:>4}/{wra:>4.0f}%/{a["pnl"]:>+11,.0f} '
          f'{b["n"]:>4}/{wrb:>4.0f}%/{b["pnl"]:>+11,.0f}')

# Specific case: 2025-04-07 Trump tariff
print()
print('=' * 80)
print('CASE STUDY: 2025-04-07 Trump tariff (same-day 2-entry sniff)')
print('=' * 80)
from datetime import date
target = date(2025, 4, 7)
print(f'\nCool=0 trades on {target}:')
for t in A['trades']:
    if t['entry_d'] == target:
        print(f'  {t["entry_sig"]} -> {t["exit_sig"]} ({t["exit_d"]}) PnL {t["pnl"]:+,.0f}')
print(f'\nCool=1 trades on {target}:')
for t in B['trades']:
    if t['entry_d'] == target:
        print(f'  {t["entry_sig"]} -> {t["exit_sig"]} ({t["exit_d"]}) PnL {t["pnl"]:+,.0f}')

# Find which trades got removed
print()
print('=' * 80)
print('TRADES IN COOL=0 BUT NOT IN COOL=1 (filtered out by Cooldown)')
print('=' * 80)
a_keys = {(t['entry_d'], t['pnl']) for t in A['trades']}
b_keys = {(t['entry_d'], t['pnl']) for t in B['trades']}
removed = sorted([(t['entry_d'], t['pnl'], t['exit_sig']) for t in A['trades']
                  if (t['entry_d'], t['pnl']) not in b_keys])
print(f'Total trades removed: {len(removed)}')
removed_pnl = sum(t[1] for t in removed)
removed_wins = sum(1 for t in removed if t[1] > 0)
print(f'Removed PnL: {removed_pnl:+,.0f}')
print(f'Removed wins/losses: {removed_wins} / {len(removed)-removed_wins}')

print('\nLargest WINNERS removed by Cooldown=1:')
for d, p, s in sorted(removed, key=lambda x: -x[1])[:10]:
    print(f'  {d}  {p:>+10,.0f}  {s}')

print('\nLargest LOSERS removed by Cooldown=1:')
for d, p, s in sorted(removed, key=lambda x: x[1])[:10]:
    print(f'  {d}  {p:>+10,.0f}  {s}')

# Net effect of removed trades
print()
print('=' * 80)
print('NET EFFECT OF COOLDOWN=1 (mechanical comparison)')
print('=' * 80)
print(f'Cool=0 Net: {A["metrics"]["淨利"]:>+,.0f}')
print(f'Cool=1 Net: {B["metrics"]["淨利"]:>+,.0f}')
print(f'Diff:       {B["metrics"]["淨利"] - A["metrics"]["淨利"]:>+,.0f}')
print(f'  → removed {len(removed)} trades, net contribution {removed_pnl:+,.0f}')
print(f'  → average per removed trade: {removed_pnl/len(removed) if removed else 0:+,.0f}')
