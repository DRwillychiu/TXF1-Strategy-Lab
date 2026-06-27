"""Analyze S3_L VolSqueezeLong WFA - 18 xlsx files (9 IS/OOS pairs)."""
import sys, io, os
from openpyxl import load_workbook
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

BASE = r'C:\Users\User\Downloads'
# 9 pairs based on filenames
pairs = [
    ('2020',       'TXF1  S3_VolSqueezeLong 策略回測績效報告IS2020.xlsx',       'TXF1  S3_VolSqueezeLong 策略回測績效報告OS2020.xlsx'),
    ('2020-2022',  'TXF1  S3_VolSqueezeLong 策略回測績效報告IS2020-2022.xlsx',  'TXF1  S3_VolSqueezeLong 策略回測績效報告OS2020-2022.xlsx'),
    ('2021-2022',  'TXF1  S3_VolSqueezeLong 策略回測績效報告IS2021-2022.xlsx',  'TXF1  S3_VolSqueezeLong 策略回測績效報告OS2021-2022.xlsx'),
    ('2021-2023',  'TXF1  S3_VolSqueezeLong 策略回測績效報告IS2021-2023.xlsx',  'TXF1  S3_VolSqueezeLong 策略回測績效報告OS2021-2023.xlsx'),
    ('2022-2023',  'TXF1  S3_VolSqueezeLong 策略回測績效報告IS2022-2023.xlsx',  'TXF1  S3_VolSqueezeLong 策略回測績效報告OS2022-2023.xlsx'),
    ('2022-2024',  'TXF1  S3_VolSqueezeLong 策略回測績效報告IS2022-2024.xlsx',  'TXF1  S3_VolSqueezeLong 策略回測績效報告OS2022-2024.xlsx'),
    ('2023-2024',  'TXF1  S3_VolSqueezeLong 策略回測績效報告IS2023-2024.xlsx',  'TXF1  S3_VolSqueezeLong 策略回測績效報告OS2023-2024.xlsx'),
    ('2023-2025',  'TXF1  S3_VolSqueezeLong 策略回測績效報告IS2023-2025.xlsx',  'TXF1  S3_VolSqueezeLong 策略回測績效報告OS2023-2025.xlsx'),
    ('2024-2025',  'TXF1  S3_VolSqueezeLong 策略回測績效報告IS2024-2025.xlsx',  'TXF1  S3_VolSqueezeLong 策略回測績效報告OS2024-2025.xlsx'),
]

def extract(path):
    """Return dict with: start, end, n, pf, sharpe, net_profit, wr, dd_pct"""
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:
        return None
    out = {}
    # 設定 sheet
    ws = wb['設定']
    for row in ws.iter_rows(values_only=True):
        if row[0] == '開始日期':
            out['start'] = row[1]
        elif row[0] == '結束日期':
            out['end'] = row[1]
    # 策略分析 sheet
    ws = wb['策略分析']
    for row in ws.iter_rows(min_row=1, max_row=80, values_only=True):
        if row[0] == '淨利':
            out['net_profit'] = row[1]
        elif row[0] == '獲利因子':
            out['pf'] = row[1]
        elif row[0] == '調整獲利因子':
            out['adj_pf'] = row[1]
        elif row[0] == '交易總次數':
            out['n'] = row[1]
        elif row[0] == '%勝率':
            out['wr'] = row[1]
        elif row[0] == '年度夏普比率':
            out['sharpe'] = row[1]
        elif row[0] == '最大策略虧損 (%)':
            out['dd_pct'] = row[1]
        elif row[0] == '最大策略虧損':
            out['dd_abs'] = row[1]
    return out

print(f'{"Window":<12} {"Phase":<4} {"Period":<24} {"N":>4} {"PF":>6} {"adjPF":>6} {"Sharpe":>7} {"NetProfit":>10} {"WR":>5} {"DD%":>7}')
print('-' * 100)

results = {}
for tag, is_file, os_file in pairs:
    is_data = extract(os.path.join(BASE, is_file))
    os_data = extract(os.path.join(BASE, os_file))
    results[tag] = (is_data, os_data)
    for phase, d in [('IS', is_data), ('OOS', os_data)]:
        if d is None:
            print(f'{tag:<12} {phase:<4} (failed to load)')
            continue
        try:
            start = d['start'].strftime('%Y-%m-%d') if d.get('start') else '?'
            end = d['end'].strftime('%Y-%m-%d') if d.get('end') else '?'
        except:
            start, end = '?', '?'
        period = f'{start[:10]}~{end[:10]}'
        n = d.get('n', '?')
        pf = d.get('pf', 0)
        adj_pf = d.get('adj_pf', 0)
        sharpe = d.get('sharpe', 0)
        np = d.get('net_profit', 0)
        wr = d.get('wr', 0)
        dd = d.get('dd_pct', 0)
        pf_s = f'{pf:.2f}' if isinstance(pf, (int, float)) else str(pf)
        adj_pf_s = f'{adj_pf:.2f}' if isinstance(adj_pf, (int, float)) else str(adj_pf)
        sharpe_s = f'{sharpe:.3f}' if isinstance(sharpe, (int, float)) else str(sharpe)
        np_s = f'{np:,.0f}' if isinstance(np, (int, float)) else str(np)
        wr_s = f'{wr:.1f}%' if isinstance(wr, (int, float)) else str(wr)
        dd_s = f'{dd:.1f}%' if isinstance(dd, (int, float)) else str(dd)
        print(f'{tag:<12} {phase:<4} {period:<24} {n:>4} {pf_s:>6} {adj_pf_s:>6} {sharpe_s:>7} {np_s:>10} {wr_s:>5} {dd_s:>7}')

# WFE calculation
print()
print('=' * 100)
print('WFE Analysis')
print('=' * 100)
print(f'{"Window":<12} {"IS PF":>7} {"OOS PF":>7} {"WFE":>7} {"OOS Sharpe":>11} {"OOS N":>6} {"Status":<20}')
print('-' * 80)

wfe_list = []
oos_pf_list = []
oos_sharpe_list = []
gate_passes = []
for tag, (is_d, os_d) in results.items():
    if is_d is None or os_d is None:
        continue
    is_pf = is_d.get('pf', 0)
    os_pf = os_d.get('pf', 0)
    os_sharpe = os_d.get('sharpe', 0)
    os_n = os_d.get('n', 0)
    if isinstance(is_pf, (int, float)) and isinstance(os_pf, (int, float)) and is_pf > 0:
        wfe = os_pf / is_pf * 100
    else:
        wfe = 0
    wfe_list.append(wfe)
    if isinstance(os_pf, (int, float)):
        oos_pf_list.append(os_pf)
    if isinstance(os_sharpe, (int, float)):
        oos_sharpe_list.append(os_sharpe)
    # 3 gates
    g1 = wfe > 50
    g2 = isinstance(os_pf, (int, float)) and os_pf > 1.0
    g3 = isinstance(os_sharpe, (int, float)) and os_sharpe > 0
    all_pass = g1 and g2 and g3
    gate_passes.append(all_pass)
    status = '✅ PASS' if all_pass else ('❌ FAIL ' + ('W' if not g1 else '') + ('P' if not g2 else '') + ('S' if not g3 else ''))
    print(f'{tag:<12} {is_pf:>7.2f} {os_pf:>7.2f} {wfe:>6.1f}% {os_sharpe:>10.3f} {os_n:>6} {status:<20}')

# Summary stats
print()
print('=' * 100)
print('SUMMARY STATS')
print('=' * 100)
n_pass = sum(gate_passes)
n_total = len(gate_passes)
mean_wfe = sum(wfe_list) / len(wfe_list) if wfe_list else 0
median_wfe = sorted(wfe_list)[len(wfe_list) // 2] if wfe_list else 0
mean_os_pf = sum(oos_pf_list) / len(oos_pf_list) if oos_pf_list else 0
mean_os_sharpe = sum(oos_sharpe_list) / len(oos_sharpe_list) if oos_sharpe_list else 0
oos_pf_pos = sum(1 for p in oos_pf_list if p > 1.0)
oos_sharpe_pos = sum(1 for s in oos_sharpe_list if s > 0)

print(f'Windows analyzed:           {n_total}')
print(f'Windows passing 3 gates:    {n_pass}/{n_total} ({n_pass/n_total*100:.1f}%)')
print(f'Mean WFE:                   {mean_wfe:.1f}%')
print(f'Median WFE:                 {median_wfe:.1f}%')
print(f'Mean OOS PF:                {mean_os_pf:.2f}')
print(f'Mean OOS Sharpe:            {mean_os_sharpe:.3f}')
print(f'OOS PF > 1.0:               {oos_pf_pos}/{len(oos_pf_list)} ({oos_pf_pos/len(oos_pf_list)*100:.1f}%)')
print(f'OOS Sharpe > 0:             {oos_sharpe_pos}/{len(oos_sharpe_list)} ({oos_sharpe_pos/len(oos_sharpe_list)*100:.1f}%)')

print()
print('=' * 100)
print('VERDICT — Walk-Forward Robustness')
print('=' * 100)
overall_gate = (n_pass / n_total >= 5/9) and (mean_wfe > 50) and (mean_os_pf > 1.0)
print(f'  ≥ 5/9 windows pass 3 gates:  {"PASS" if n_pass/n_total >= 5/9 else "FAIL"}  ({n_pass}/{n_total})')
print(f'  Mean WFE > 50%:              {"PASS" if mean_wfe > 50 else "FAIL"}  ({mean_wfe:.1f}%)')
print(f'  Mean OOS PF > 1.0:           {"PASS" if mean_os_pf > 1.0 else "FAIL"}  ({mean_os_pf:.2f})')
print()
if overall_gate:
    print('  ✅ S3_L WFA PASSES → eligible for W5 10-dim eval')
else:
    print('  ❌ S3_L WFA FAILS → likely overfit, recommend退回 to Phase 1 baseline or KILL')
